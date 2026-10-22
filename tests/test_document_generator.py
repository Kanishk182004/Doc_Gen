import pytest
from pathlib import Path
from docx import Document
from app.schemas import DocumentContent, SectionContent
from app.tools.document_generator import create_docx_file

def test_create_docx_file(tmp_path: Path):
    doc_content = DocumentContent(
        title="Test Document Title",
        subtitle="Test Subtitle",
        document_type="Project Plan",
        executive_summary="This is an executive summary test for docx generator.",
        sections=[
            SectionContent(
                heading="1. Section Heading",
                level=1,
                body="This is paragraph body text.",
                bullet_points=["Bullet 1", "Bullet 2"],
                table_headers=["Col A", "Col B"],
                table_rows=[["Val 1", "Val 2"]]
            )
        ],
        assumptions=["Assumption 1"],
        next_steps=["Step 1"]
    )

    out_file = tmp_path / "test_doc.docx"
    result_path = create_docx_file(doc_content, out_file)

    assert result_path.exists()
    assert result_path.stat().st_size > 0

    # Verify docx structure using python-docx
    doc = Document(str(result_path))
    headings = [p.text for p in doc.paragraphs if p.text]
    assert any("Test Document Title" in h for h in headings)
    assert any("1. Section Heading" in h for h in headings)
    assert len(doc.tables) >= 2  # Metadata table + data table
